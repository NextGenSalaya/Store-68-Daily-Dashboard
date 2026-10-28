"""Topic 4 of the Store 68 Daily Board: Over Stock items and open orders on the floor (On order Floor).
Reads the two source files from anywhere under this folder and writes JSON datasets.
Usage: python build_topic4.py [out_folder]
"""
import glob, json, os, re, sys, warnings
from datetime import datetime, date
import openpyxl

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "data_t4")
os.makedirs(OUT, exist_ok=True)


def find(pattern):
    hits = sorted(glob.glob(os.path.join(HERE, "**", pattern), recursive=True), key=os.path.getmtime)
    if not hits:
        sys.exit(f"{pattern} not found under {HERE}")
    print("reading", hits[-1])
    return hits[-1]


def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def rnd(v, d=2):
    v = num(v)
    return None if v is None else round(v, d)


def txt(v):
    return "" if v is None else str(v).strip()


def code(v):
    n = num(v)
    return str(int(n)) if n is not None and n == int(n) else txt(v)


def div_of(dept):
    n = num(dept) or 0
    return "FRESH FOOD" if n <= 5 or n in (19, 21) else "DRY FOOD" if n <= 17 or n == 22 else "NON FOOD"


def save(name, rows):
    with open(os.path.join(OUT, name + ".json"), "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, separators=(",", ":"))
    print(f"{name}: {len(rows)} rows")


meta = []

# ---------- OVER STOCK: one block per division, each with its own header row ----------
path = find("OVER STOCK*.xlsx")
ws = openpyxl.load_workbook(path, read_only=True, data_only=True).active
over, div = {}, ""
for r in ws.iter_rows(values_only=True):
    a = txt(r[0])
    if a.upper() in ("DRY FOOD", "FRESH FOOD", "NON FOOD"):
        div = a.upper()
        continue
    if num(r[0]) is None or not txt(r[2]):
        continue
    item = code(r[2])
    over[item] = {"key": item, "division": div or div_of(r[0]), "dept": code(r[0]), "cls": code(r[1]), "item": item,
                  "desc": txt(r[3]), "soh": rnd(r[4]), "retail": rnd(r[5]), "way": txt(r[6]) or "(ว่าง)",
                  "on_order": rnd(r[7]), "dms": rnd(r[8]), "stock_days": rnd(r[9], 1)}
save("t4_over", list(over.values()))
meta.append({"key": "Over Stock file", "value": datetime.fromtimestamp(os.path.getmtime(path)).strftime("%Y-%m-%d")})

# ---------- On order Floor: open order lines with planned delivery ----------
ws = openpyxl.load_workbook(find("On order Floor*.xlsx"), read_only=True, data_only=True).active
m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})\s+(\d{1,2}:\d{2})", txt(ws["A2"].value))
run = f"{m.group(3)}-{int(m.group(1)):02d}-{int(m.group(2)):02d} {m.group(4)}" if m else ""
meta.append({"key": "On order run", "value": run})
floor = {}
for r in ws.iter_rows(min_row=5, values_only=True):
    if not txt(r[5]) or num(r[3]) is None:
        continue
    d = r[0].strftime("%Y-%m-%d") if isinstance(r[0], (datetime, date)) else txt(r[0])[:10]
    k = f"{code(r[1])}|{code(r[5])}"
    if k in floor:
        floor[k]["qty"] = round(floor[k]["qty"] + (num(r[7]) or 0), 2)
        continue
    floor[k] = {"key": k, "plan_date": d, "order": code(r[1]), "dept": code(r[3]), "division": div_of(r[3]), "cls": code(r[4]),
                "item": code(r[5]), "desc": txt(r[6]), "qty": rnd(r[7]) or 0, "status": txt(r[8]), "type": txt(r[9]),
                "plan": txt(r[10])}
save("t4_floor", list(floor.values()))
save("t4_meta", meta)

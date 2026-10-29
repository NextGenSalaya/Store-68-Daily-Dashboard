"""Read every data sheet of "All Dashboard Store68.xlsx" into JSON datasets for the Store 68 Daily Board.

The file is found anywhere under this folder, so it can sit in a sub-folder.
Usage: python build_all_dashboard.py [out_folder]
"""
import glob, json, os, re, sys, warnings
from datetime import datetime, date
import openpyxl
from openpyxl.utils import column_index_from_string as ci, get_column_letter as cl

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "data_ad")
os.makedirs(OUT, exist_ok=True)

found = sorted(glob.glob(os.path.join(HERE, "**", "All Dashboard Store68*.xlsx"), recursive=True), key=os.path.getmtime)
if not found:
    sys.exit("All Dashboard Store68.xlsx not found under " + HERE)
SRC = found[-1]
print("reading", SRC)
wb = openpyxl.load_workbook(SRC, data_only=True)


def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("%", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def rnd(v, d=2):
    v = num(v)
    return None if v is None else round(v, d)


def txt(v):
    return "" if v is None else str(v).strip()


def iso(v):
    return v.strftime("%Y-%m-%d") if isinstance(v, (datetime, date)) else None


def code(v):
    """Item / class numbers as text without a trailing .0"""
    n = num(v)
    return str(int(n)) if n is not None and n == int(n) else txt(v)


def save(name, rows):
    with open(os.path.join(OUT, name + ".json"), "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, separators=(",", ":"))
    print(f"{name}: {len(rows)} rows")


def header(ws, row, c1, c2):
    return [(c, txt(ws.cell(row, c).value)) for c in range(ci(c1), ci(c2) + 1)]


# ------------------------------------------------------------------ Dashboard sheet
ws = wb["Dashboard"]
m = re.search(r"(\d{2})-(\d{2})-(\d{4})", txt(ws["A3"].value))
DATA_DATE = f"{m.group(3)}-{m.group(2)}-{m.group(1)}" if m else None
PERIODS = ["Daily", "WTD", "MTD", "YTD"]


def period_table(hrow, fields):
    cols = [c for c, t in header(ws, hrow, "B", "CH") if t]
    n = len(fields)
    out, r = [], hrow + 1
    while txt(ws.cell(r, 1).value):
        div = txt(ws.cell(r, 1).value)
        for p, per in enumerate(PERIODS):
            rec = {"key": f"{per}|{div}", "period": per, "division": div}
            for c, f in zip(cols[p * n:(p + 1) * n], fields):
                rec[f] = rnd(ws.cell(r, c).value)
            out.append(rec)
        if div.lower().startswith("total"):
            break
        r += 1
    return out


save("sales", period_table(6, ["actual", "budget", "pct_budget", "ly", "pct_ly", "share"]))
save("margin", period_table(21, ["actual", "budget", "var_budget", "ly", "var_ly", "share"]))

save("daily_need", [{"division": "Total" if txt(ws[f"A{r}"].value) == "Grand Total" else txt(ws[f"A{r}"].value),
                     "avg_day": rnd(ws[f"C{r}"].value), "budget_day": rnd(ws[f"G{r}"].value),
                     "pct_var": rnd(ws[f"K{r}"].value), "need_day": rnd(ws[f"Q{r}"].value)}
                    for r in range(14, 18) if txt(ws[f"A{r}"].value)])

mix = []
for r in range(14, 18):
    div = txt(ws[f"X{r}"].value)
    for kind, cols in {"Makro Mail": "AD AH AL", "Own Brand": "AT AZ BG", "KVI": "BN BX CC"}.items():
        d, mt, y = cols.split()
        mix.append({"key": f"{div}|{kind}", "division": div, "kind": kind,
                    "daily": rnd(ws[f"{d}{r}"].value), "mtd": rnd(ws[f"{mt}{r}"].value), "ytd": rnd(ws[f"{y}{r}"].value)})
save("mix", mix)

hdr = [c for c, t in header(ws, 29, "B", "CH") if t]
cf = ["invoices", "invoices_sply", "sales", "sales_sply", "active", "active_sply", "visits", "visits_sply", "participation"]
cust, r = [], 30
while txt(ws.cell(r, 1).value):
    g = txt(ws.cell(r, 1).value)
    for p, per in enumerate(["WTD", "MTD"]):
        rec = {"key": f"{per}|{g}", "period": per, "group": g}
        for c, f in zip(hdr[p * 9:(p + 1) * 9], cf):
            rec[f] = rnd(ws.cell(r, c).value)
        cust.append(rec)
    r += 1
save("customers", cust)

kpi = [{"metric": txt(ws[f"A{r}"].value), "value": rnd(ws[f"D{r}"].value)} for r in range(39, 48) if txt(ws[f"A{r}"].value)]
kpi.append({"metric": "Data date", "value": DATA_DATE})
save("kpi", kpi)

save("aging_status", [{"status": txt(ws[f"N{r}"].value), "stock_k": rnd(ws[f"S{r}"].value, 0), "items": rnd(ws[f"V{r}"].value, 0),
                       "pct_stock": rnd(ws[f"Z{r}"].value)} for r in range(54, 60) if txt(ws[f"N{r}"].value)])

save("bottom_stores", [{"store": txt(ws[f"A{r}"].value), "variance": rnd(ws[f"D{r}"].value), "pct_variance": rnd(ws[f"H{r}"].value)}
                       for r in range(54, 64) if txt(ws[f"A{r}"].value)])

ranks = []
for lst, first, cols, fields in (
        ("Top 10 by sales", 40, "AC AK AO AR AW AZ", ["mtd", "budget", "pct_budget", "ly", "pct_ly"]),
        ("Bottom 10 by sales", 40, "BC BQ BV BY CC", ["mtd", "var_ly", "pct_ly", "pct_budget"]),
        ("Top 10 by margin", 54, "AC AK AO AR AW AZ", ["mtd", "budget", "pct_budget", "ly", "pct_ly"]),
        ("Bottom 10 by margin", 54, "BC BQ BV BY CC", ["mtd", "var_ly", "pct_ly", "pct_budget"])):
    cols = cols.split()
    for i, r in enumerate(range(first, first + 10)):
        if not txt(ws[f"{cols[0]}{r}"].value):
            continue
        rec = {"key": f"{lst}|{i + 1}", "list": lst, "rank": i + 1, "class": txt(ws[f"{cols[0]}{r}"].value)}
        for c, f in zip(cols[1:], fields):
            rec[f] = rnd(ws[f"{c}{r}"].value)
        ranks.append(rec)
save("class_rank", ranks)

# ------------------------------------------------------------------ Category sheet
ws = wb["Category"]
cats, div = [], ""
for r in range(4, ws.max_row + 1):
    a, b = txt(ws[f"A{r}"].value), txt(ws[f"B{r}"].value)
    is_total = a.lower().startswith("total")
    if not b and not is_total:
        continue
    if a and not is_total:
        div = a
    name = "รวม " + a.replace("Total ", "") if is_total else b
    rec = {"key": f"{div if not is_total else a}|{name}", "division": div if not is_total else a.replace("Total ", "").replace("Total", "ทั้งร้าน"),
           "category": name, "is_total": is_total}
    for p, cols in {"day": "C D G", "mtd": "J K N", "ytd": "Q R U"}.items():
        a1, b1, l1 = cols.split()
        rec[f"{p}_actual"], rec[f"{p}_budget"], rec[f"{p}_ly"] = rnd(ws[f"{a1}{r}"].value), rnd(ws[f"{b1}{r}"].value), rnd(ws[f"{l1}{r}"].value)
    if is_total or any(abs(rec[k] or 0) >= 0.5 for k in ("mtd_actual", "mtd_budget", "ytd_actual")):
        cats.append(rec)
# classes in no category (from ST.68 Sales By Category, its "Check Diff") form a group per store division and join the
# store total, so the Category total equals the store's sales
cat_files = sorted(glob.glob(os.path.join(HERE, "**", "ST.68 Sales By Category.xlsx"), recursive=True), key=os.path.getmtime)
if cat_files:
    from unmapped import unmapped_classes, GROUP, LABEL
    extra = unmapped_classes(cat_files[-1])
    store = next((c for c in cats if c["division"] == "ทั้งร้าน"), None)
    at = cats.index(store) if store else len(cats)
    for sd in sorted({x["store_div"] for x in extra}):
        grp = f"{GROUP} · {dict(zip(('DRY FOOD', 'FRESH FOOD', 'NON FOOD'), ('DF', 'FF', 'NF'))).get(sd, sd)}"
        rows = [x for x in extra if x["store_div"] == sd]
        vals = {f"{p}_{k}": round(sum(x[f"{q}_{s}"] for x in rows), 2)
                for p, q in (("day", "d"), ("mtd", "m"), ("ytd", "y")) for k, s in (("actual", "a"), ("budget", "b"), ("ly", "ly"))}
        member = {"key": f"{grp}|{LABEL}", "division": grp, "store_div": sd, "category": f"{LABEL} ({len(rows)} Class)", "is_total": False, **vals}
        total = {"key": f"Total {grp}|รวม {grp}", "division": grp, "store_div": sd, "category": f"รวม {grp}", "is_total": True, **vals}
        cats[at:at] = [member, total]
        at += 2
        if store:
            for k, v in vals.items():
                store[k] = round((store[k] or 0) + v, 2)
    print(f"category: added {len(extra)} classes with no category (MTD {round(sum(x['m_a'] for x in extra), 2)} thousand baht)")
save("category", cats)

# ------------------------------------------------------------------ Analysis sheet (sales since 1 June and YTD by department)
ws = wb["Analysis"]
dept = []
for r in range(4, ws.max_row + 1):
    a, b = txt(ws[f"A{r}"].value), txt(ws[f"B{r}"].value)
    if not a:
        continue
    total = not b
    rec = {"key": a if total else f"{a} {b}", "dept": a if total else f"{a} {b}", "is_total": total,
           "p_actual": rnd(ws[f"C{r}"].value), "p_budget": rnd(ws[f"D{r}"].value), "p_ly": rnd(ws[f"G{r}"].value),
           "ytd_actual": rnd(ws[f"J{r}"].value), "ytd_budget": rnd(ws[f"K{r}"].value), "ytd_ly": rnd(ws[f"N{r}"].value)}
    if total or any(abs(rec[k] or 0) >= 1 for k in ("p_actual", "ytd_actual")):
        dept.append(rec)
# departments come before their division's total row
pending = []
for rec in dept:
    if rec["is_total"]:
        div = "ALL" if rec["dept"].upper().startswith("GRAND") else rec["dept"].replace(" Total", "").strip()
        rec["division"] = div
        for d in pending:
            d["division"] = div
        pending = []
    else:
        pending.append(rec)
save("dept_sales", dept)

# ------------------------------------------------------------------ RA_DATA blocks (clean source tables, header on row 4)
ra = wb["RA_DATA"]


def split_total(div):
    """'DRY FOOD Total' -> ('DRY FOOD', True)"""
    return (div[:-6].strip(), True) if div.endswith("Total") else (div, False)


def ra_block(c1, c2, keycol, carry=True):
    cols = header(ra, 4, c1, c2)
    rows, last = [], ""
    for r in range(5, ra.max_row + 1):
        vals = {t: ra.cell(r, c).value for c, t in cols}
        if ra.cell(r, ci(keycol)).value in (None, "", 0):
            if all(v in (None, "") for v in vals.values()):
                break
            continue
        first = cols[0][1]
        if carry:
            if txt(vals[first]):
                last = txt(vals[first])
            vals[first] = last
        rows.append(vals)
    return rows


rows = []
for r in ra_block("A", "H", "B"):
    if not num(r["Rank"]):
        continue
    div, tot = split_total(r["Division"])
    rows.append({"key": f"{div}|{'Total' if tot else int(num(r['Rank']))}", "division": div, "is_total": tot,
                 "rank": None if tot else int(num(r["Rank"])), "class": "รวม Top 5" if tot else f"{code(r['Class Number'])} {txt(r['Class'])}",
                 "sales": rnd(r["Net Sales Amt"], 0), "qty": rnd(r["Net Sales Qty"], 0), "profit": rnd(r["Net Profit"], 0),
                 "profit_pct": rnd(r["Profit %"])})
save("top_class", rows)

ly = {code(r["Item Number"]): rnd(r[[k for k in r if k.startswith("Net Sales")][0]], 0) for r in ra_block("BD", "BG", "BE")}
rows = []
for r in ra_block("J", "U", "N"):
    if not num(r["Rank"]):
        continue
    div, tot = split_total(r["Division"])
    rows.append({"key": f"{div}|{'Total' if tot else int(num(r['Rank']))}", "division": div, "is_total": tot,
                 "rank": None if tot else int(num(r["Rank"])), "item": "" if tot else code(r["Item Number"]),
                 "desc": "รวม Top 10" if tot else txt(r["Item"]), "class": txt(r["Class"]),
                 "qty": rnd(r["Net Sales Qty"], 1), "sales": rnd(r["Net Sales Amt"], 0), "ly": None if tot else ly.get(code(r["Item Number"])),
                 "profit": rnd(r["Net Profit"], 0), "profit_pct": rnd(r["Profit %"]), "stock": rnd(r["Stock"], 1)})
save("top_items", rows)

save("stock_div", [{"division": re.sub(r"^\d\.", "", txt(r["DIVISION"])), "sales": rnd(r["Net Sales Amt"], 0),
                    "avg_sales": rnd(r["Avg Sales Amt"], 0), "stock": rnd(r["Stock Amount"], 0), "stock_day": rnd(r["Stock Day"], 1)}
                   for r in ra_block("W", "AA", "W", carry=False) if num(r["Stock Amount"])])

save("aging_items", [{"key": f"{r['Division']}|{int(num(r['RANK']))}", "division": r["Division"], "rank": int(num(r["RANK"])),
                      "status": txt(r["Status"]), "dept": code(r["Dept"]), "item": code(r["Item Number"]), "desc": txt(r["Item"]),
                      "price": rnd(r["Price"]), "qty": rnd(r["EOH Qty"], 1), "cost": rnd(r["EOH Cost"], 0),
                      "last_sold": iso(r["Last Sold Date"]), "last_received": iso(r["Last Received Date"])}
                     for r in ra_block("AC", "AN", "AH") if num(r["RANK"])])

rows = []
for r in [dict(zip([t for _, t in header(ra, 1, "BI", "BQ")], [ra.cell(rr, c).value for c in range(ci("BI"), ci("BQ") + 1)]))
          for rr in range(2, ra.max_row + 1)]:
    if not num(r["TopShk."]):
        continue
    tot = txt(r["Dept.No"]).startswith("Grand Total")
    rows.append({"key": "Total" if tot else str(int(num(r["TopShk."]))), "rank": None if tot else int(num(r["TopShk."])),
                 "is_total": tot, "dept": "รวม" if tot else f"{code(r['Dept.No'])} {txt(r['Department'])}",
                 "class": "" if tot else f"{code(r['Class Number'])} {txt(r['Class'])}", "item": code(r["Item Number"]),
                 "desc": txt(r["Item"]), "qty": rnd(r["Shrinkage Qty"], 1), "cost": rnd(r["Shrinkage Cost Amt"], 0)})
shrink_rows = rows  # division is added once the Shk sheet has been read

profit = []
for lst, c1, c2 in (("กำไรสูงสุด", "BS", "CB"), ("ขาดทุนมากสุด", "CD", "CM")):
    for r in ra_block(c1, c2, cl(ci(c1) + 4)):
        if not num(r["Rank"]):
            continue
        profit.append({"key": f"{lst}|{r['Division']}|{int(num(r['Rank']))}", "list": lst, "division": r["Division"],
                       "rank": int(num(r["Rank"])), "class": f"{code(r['Class Number'])} {txt(r['Class'])}", "item": code(r["Item Number"]),
                       "desc": txt(r["Item"]), "qty": rnd(r["Net Sales Qty"], 1), "sales": rnd(r["Net Sales Amt"], 0),
                       "profit": rnd(r["Net Profit"], 0), "profit_pct": rnd(r["Profit %"])})
save("profit_items", profit)

# ------------------------------------------------------------------ Shrinkage: by reason code (CodeID) and vs budget (Shk)
ws = wb["CodeID"]
codes = []
for per, first in (("MTD", 6), ("YTD", 32)):
    for r in range(first, first + 20):
        a, b = txt(ws[f"A{r}"].value), txt(ws[f"B{r}"].value)
        if not a:
            continue
        name = "TOTAL" if a == "TOTAL" else f"{a} {b}"
        codes.append({"key": f"{per}|{name}", "period": per, "dept": name, "adjust_minus": rnd(ws[f"C{r}"].value, 0),
                      "adjust_plus": rnd(ws[f"D{r}"].value, 0), "joblot": rnd(ws[f"E{r}"].value, 0), "expired": rnd(ws[f"F{r}"].value, 0),
                      "stock_count": rnd(ws[f"G{r}"].value, 0), "markdown": rnd(ws[f"H{r}"].value, 0), "total": rnd(ws[f"I{r}"].value, 0),
                      "sales": rnd(ws[f"J{r}"].value, 0), "pct_sales": rnd((num(ws[f"K{r}"].value) or 0) * 100, 3)})
        if a == "TOTAL":
            break

ws = wb["Shk"]
# department names: the Analysis sheet lists every department, CodeID only those with shrinkage
names = {d["dept"].split(" ", 1)[0]: d["dept"] for d in dept if not d["is_total"]}
names.update({c["dept"].split(" ", 1)[0]: c["dept"] for c in codes if c["dept"] != "TOTAL"})
# The Shk sheet lists departments in division blocks closed by a total row; that gives each department its division
TOTAL_DIV = {"Total Fresh": "FRESH FOOD", "Total Dry Food": "DRY FOOD", "Total Non Food": "NON FOOD", "Total Store": "ALL"}
block, dept_div = [], {}
for r in range(4, ws.max_row + 1):
    e = txt(ws[f"E{r}"].value)
    if e in TOTAL_DIV:
        for d in block:
            dept_div[d] = TOTAL_DIV[e]
        block = []
    elif e:
        block.append(e)
shk = []
for r in range(4, ws.max_row + 1):
    e = txt(ws[f"E{r}"].value)
    if not e or num(ws[f"F{r}"].value) is None:
        continue
    for per, cols in (("MTD", "F G H I"), ("YTD", "J K L M")):
        s, k, p, b = cols.split()
        shk.append({"key": f"{per}|{e}", "period": per, "dept": names.get(e, e), "is_total": e.startswith("Total"),
                    "division": TOTAL_DIV.get(e) or dept_div.get(e, ""),
                    "sales": rnd(ws[f"{s}{r}"].value, 0), "shrink": rnd(ws[f"{k}{r}"].value, 0),
                    "pct_sales": rnd(ws[f"{p}{r}"].value, 3), "pct_budget": rnd(ws[f"{b}{r}"].value, 3)})
save("shrink_budget", shk)
def div_of(dept_no):
    """Division of a department number; departments missing from the Shk sheet follow the store's numbering."""
    if dept_no in dept_div:
        return dept_div[dept_no]
    n = num(dept_no) or 0
    return "FRESH FOOD" if n <= 5 or n in (19, 21) else "DRY FOOD" if n <= 17 or n == 22 else "NON FOOD"


for c in codes:
    c["division"] = "ALL" if c["dept"] == "TOTAL" else div_of(c["dept"].split(" ", 1)[0])
save("shrink_code", codes)
for s in shrink_rows:
    s["division"] = "ALL" if s["is_total"] else div_of(s["dept"].split(" ", 1)[0])
save("shrink_items", shrink_rows)

# ------------------------------------------------------------------ Service level: KVI, Makro Mail (KVI-mm) and Top 100 (tOP100)
ws = wb["KVI-mm"]
svc = []
for lst, first in (("KVI", 5), ("Makro Mail", 13)):
    for c, div in (("D", "FF"), ("E", "DF"), ("F", "NF"), ("G", "Total")):
        svc.append({"key": f"{lst}|{div}", "list": lst, "division": div, "in_stock": rnd(ws[f"{c}{first}"].value, 0),
                    "oos_not_on_order": rnd(ws[f"{c}{first + 1}"].value, 0), "oos_on_order": rnd(ws[f"{c}{first + 2}"].value, 0),
                    "total": rnd(ws[f"{c}{first + 3}"].value, 0), "service_level": rnd(ws[f"{c}{first + 4}"].value)})
save("service_level", svc)
mm = re.search(r"#\S+.*", txt(ws["B12"].value))
save("service_note", [{"list": "Makro Mail", "note": txt(ws["B12"].value)}])

ws = wb["tOP100"]
fields = ["total", "stock_pos", "stock_zero", "stock_neg", "oos", "oos_on_order", "oos_not_on_order", "oos_info3", "not_order_store", "not_order_ho"]
t100 = []
for c in range(ci("C"), ci("O") + 1):
    sec = txt(ws.cell(3, c).value)
    if not sec:
        continue
    div = {"TT_Fresh": "FRESH FOOD", "TT_DF": "DRY FOOD", "TT_NF": "NON FOOD", "Total_All": "ALL"}.get(sec) or (
        "DRY FOOD" if sec.startswith("DF") else "NON FOOD" if sec.startswith("NF") else "FRESH FOOD")
    rec = {"section": sec, "division": div, "is_total": sec.startswith("TT") or sec.startswith("Total")}
    for f, r in zip(fields, range(4, 14)):
        rec[f] = rnd(ws.cell(r, c).value, 0)
    rec["service_level"] = rnd((num(ws.cell(14, c).value) or 0) * 100)
    t100.append(rec)
save("top100", t100)
save("gap_loss", [{"week": txt(ws[f"R{r}"].value), "gap_round1": rnd(ws[f"S{r}"].value, 0), "gap_round2": rnd(ws[f"T{r}"].value, 0),
                   "gap_total": rnd(ws[f"U{r}"].value, 0), "loss_sale": rnd(ws[f"V{r}"].value, 0)}
                  for r in range(4, 10) if txt(ws[f"R{r}"].value)])

# ------------------------------------------------------------------ Top 4000 stock status by reorder way (store 68 block)
ws = wb["Top4000"]
t4, status = [], ""
start = next(r for r in range(1, ws.max_row + 1) if txt(ws[f"A{r}"].value) == "Store")
for r in range(start + 1, ws.max_row + 1):
    b, c = txt(ws[f"B{r}"].value), txt(ws[f"C{r}"].value)
    a = txt(ws[f"A{r}"].value)
    if a.endswith("Total"):
        label, total = "รวมทั้งหมด", True
    elif b.startswith("Total"):
        label, total = b, True
    elif c:
        status = b or status
        label, total = f"{status} · {c}", False
    else:
        continue
    t4.append({"key": label, "status": label, "is_total": total,
               "sales_df": rnd(ws[f"D{r}"].value, 0), "sales_ff": rnd(ws[f"E{r}"].value, 0), "sales_nf": rnd(ws[f"F{r}"].value, 0),
               "sales": rnd(ws[f"G{r}"].value, 0), "sku_df": rnd(ws[f"H{r}"].value, 0), "sku_ff": rnd(ws[f"I{r}"].value, 0),
               "sku_nf": rnd(ws[f"J{r}"].value, 0), "skus": rnd(ws[f"K{r}"].value, 0), "part": rnd((num(ws[f"O{r}"].value) or 0) * 100)})
save("top4000", t4)

# ------------------------------------------------------------------ O2O channels
ws = wb["O2O"]
o2o = []
for r in range(4, ws.max_row + 1):
    a, b = txt(ws[f"A{r}"].value), txt(ws[f"B{r}"].value)
    if not a:
        continue
    total = a == "Grand Total"
    rec = {"key": "Grand Total" if total else f"{a} · {b}", "channel": a, "sub": "" if total else b, "is_total": total,
           "day_inv": rnd(ws[f"D{r}"].value, 0), "day_sales": rnd(ws[f"E{r}"].value), "day_inv_ly": rnd(ws[f"F{r}"].value, 0),
           "day_sales_ly": rnd(ws[f"G{r}"].value), "mtd_inv": rnd(ws[f"H{r}"].value, 0), "mtd_sales": rnd(ws[f"I{r}"].value),
           "mtd_inv_ly": rnd(ws[f"J{r}"].value, 0), "mtd_sales_ly": rnd(ws[f"K{r}"].value)}
    o2o.append(rec)
save("o2o", o2o)
o2o_share = [{"metric": "O2O % of store sales", "daily": rnd(ws["F26"].value), "mtd": rnd(ws["J26"].value)}]
save("o2o_share", o2o_share)

print("data date:", DATA_DATE)

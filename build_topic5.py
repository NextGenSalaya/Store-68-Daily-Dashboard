"""Topic 5 of the Store 68 Daily Board: stock days, stock MU, aging & NBS, Top 1000 / 4000 / 50 FF, sales by category,
sale by item, top drop customers, shrinkage by code, last sales 3 days, negative stock, gone forever & discontinue,
hamper and soft-serve / rotisserie sales. Reads the 19 source files from anywhere under this folder.

Customer names and numbers (Top drop Customer) and user ids / comments (Negative stock, Adj sheet) are never written
out: customers are only counted and summed. Sheets whose run date is older than the rest (Top Aging: Item, sum;
Top1000: store 68, FV..NF2; Top_4000: Data2 with #REF!) are skipped. Percent values are in percent units.
Usage: python build_topic5.py [out_folder]
"""
import glob, json, os, re, sys, warnings
from collections import defaultdict
from datetime import datetime, date, timedelta
import openpyxl

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "data_t5")
os.makedirs(OUT, exist_ok=True)
meta = []


def find(pattern):
    hits = sorted(glob.glob(os.path.join(HERE, "**", pattern), recursive=True), key=os.path.getmtime)
    if not hits:
        sys.exit(f"{pattern} not found under {HERE}")
    print("reading", hits[-1])
    return hits[-1]


def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("%", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def rnd(v, d=2):
    v = num(v)
    return None if v is None else round(v, d)


def pct(v, d=2):
    v = num(v)
    return None if v is None else round(v * 100, d)


def txt(v):
    return "" if v is None else str(v).strip()


def code(v):
    n = num(v)
    return str(int(n)) if n is not None and n == int(n) else txt(v)


def day(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)) and 30000 < v < 60000:
        return (datetime(1899, 12, 30) + timedelta(days=int(v))).strftime("%Y-%m-%d")
    s = txt(v)
    m = re.match(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", s)
    return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}" if m else None


def div_of(dept):
    n = num(dept) or 0
    return "FRESH FOOD" if n <= 5 or n in (19, 21) else "DRY FOOD" if n <= 17 or n == 22 else "NON FOOD"


DIVN = {"FF": "FRESH FOOD", "DF": "DRY FOOD", "NF": "NON FOOD", "F": "FRESH FOOD", "D": "DRY FOOD"}


def save(name, rows):
    with open(os.path.join(OUT, name + ".json"), "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, separators=(",", ":"))
    print(f"{name}: {len(rows)} rows")


def sheet(f, s=None):
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    return wb[s] if s else wb.worksheets[0]


def xlsb_rows(f, s):
    from pyxlsb import open_workbook
    with open_workbook(f) as wb:
        with wb.get_sheet(s) as ws:
            return [[c.v for c in r] for r in ws.rows()]


def dedupe(rows):
    out = {}
    for r in rows:
        out.setdefault(r["key"], r)
    return list(out.values())


# ---------- Stock Days ----------
ws = sheet(find("Stock Days by division.xlsx"))
meta.append({"key": "Stock Days date", "value": day(ws["B3"].value)})
save("t5_sd_div", [{"key": "ALL" if txt(r[0]) == "Grand Total" else txt(r[0]), "division": "ALL" if txt(r[0]) == "Grand Total" else txt(r[0]),
                    "amount": rnd(r[1]), "days": rnd(r[2], 1)} for r in ws.iter_rows(min_row=5, max_col=3, values_only=True) if txt(r[0])])
ws = sheet(find("Stock Days by division by dept.xlsx"))
save("t5_sd_dept", [{"key": code(r[3]), "division": txt(r[0]), "dept": code(r[3]), "amount": rnd(r[1]), "days": rnd(r[2], 1)}
                    for r in ws.iter_rows(min_row=5, max_col=4, values_only=True) if num(r[3]) is not None])

ws = sheet(find("Stock Last Received Date 30 Day.xlsx"))
recv, dv = [], ""
for r in ws.iter_rows(min_row=1, max_col=7, values_only=True):
    if txt(r[0]).upper() in ("FRESH FOOD", "DRY FOOD", "NON FOOD"):
        dv = txt(r[0]).upper(); continue
    if num(r[0]) is None or not txt(r[2]):
        continue
    recv.append({"key": code(r[2]), "division": dv or div_of(r[0]), "dept": code(r[0]), "cls": code(r[1]), "item": code(r[2]),
                 "desc": txt(r[3]), "last_recv": day(r[4]), "soh": rnd(r[5]), "cost": rnd(r[6])})
save("t5_recv30", dedupe(recv))

# ---------- Stock 1-3 MU / 1-5 MU ----------
ws = sheet(find("Stock1-3MU_Report.xlsx"))
mu = []
for r in ws.iter_rows(min_row=4, max_col=19, values_only=True):
    if not txt(r[6]):
        continue
    mu.append({"key": code(r[6]), "buyer": txt(r[0]), "division": txt(r[1]).upper(), "dept": code(r[2]), "dept_name": txt(r[3]),
               "cls": code(r[4]), "cls_name": txt(r[5]), "item": code(r[6]), "desc": txt(r[7]) or txt(r[8]), "last_recv": day(r[12]),
               "last_sold": day(r[13]), "aging": txt(r[14]), "on_order": rnd(r[15]), "cover": rnd(r[16], 0), "qty": rnd(r[17]), "cost": rnd(r[18])})
save("t5_mu13", dedupe(mu))

import xlrd
wb = xlrd.open_workbook(find("Stock1-5 MU_Report.xls"), encoding_override="cp874")
mu5 = []
for s in wb.sheets():
    for i in range(1, s.nrows):
        r = s.row_values(i)
        if num(r[1]) is None:
            continue
        mu5.append({"key": f"{s.name}|{code(r[1])}", "division": DIVN.get(s.name, div_of(r[0])), "dept": code(r[0]), "item": code(r[1]),
                    "cls": code(r[2]), "desc": txt(r[3]), "price": rnd(r[5]), "stock": rnd(r[6]), "amount": rnd(r[7]), "on_order": rnd(r[8])})
save("t5_mu15", dedupe(mu5))

# ---------- Top Aging & NBS (current sheets only) ----------
f = find("Top Aging&NBS by Division_V4.xlsb")
SUBS = ["FV", "BU", "FS", "BK", "FZ", "FF", "DF1", "DF2", "DF", "NF", "Total"]
trend = []
for r in xlsb_rows(f, "Summary")[3:]:
    if len(r) < 26 or day(r[1]) is None:
        continue
    rec = {"key": day(r[1]), "date": day(r[1])}
    for i, s in enumerate(SUBS):
        rec["nbs_" + s] = rnd(r[2 + i]); rec["ag_" + s] = rnd(r[15 + i])
    if rec["nbs_Total"] or rec["ag_Total"]:
        trend.append(rec)
save("t5_ag_trend", trend)
ag = []
for r in xlsb_rows(f, "RA_Aging")[1:]:
    r = r + [None] * 26
    if num(r[10]) is None:
        continue
    costs = [num(x) or 0 for x in r[12:17]]
    qtys = [num(x) or 0 for x in r[18:23]]
    ag.append({"key": code(r[10]), "division": txt(r[8]).upper(), "status": txt(r[9]), "item": code(r[10]), "desc": txt(r[11]),
               "cost": round(sum(costs), 2), "qty": round(sum(qtys), 2), "dept": code(r[17]), "cls": code(r[25]) if num(r[25]) else ""})
save("t5_aging_items", dedupe(ag))
nb = []
for r in xlsb_rows(f, "RA_NBS")[1:]:
    r = r + [None] * 24
    if num(r[9]) is None:
        continue
    nb.append({"key": code(r[9]), "division": txt(r[7]).upper(), "provision": code(r[8]), "item": code(r[9]), "desc": txt(r[10]),
               "cost": round(sum(num(x) or 0 for x in r[11:15]), 2), "dept": code(r[15]), "qty": round(sum(num(x) or 0 for x in r[16:20]), 2),
               "cls": code(r[23]) if num(r[23]) else ""})
save("t5_nbs_items", dedupe(nb))

# ---------- Top 4000 / Top 1000 ----------
f = find("Top_4000.xlsx")
ws = sheet(f, "Summary")
meta.append({"key": "Top 4000 date", "value": day(ws["B4"].value)})
t4s = []
for r in ws.iter_rows(min_row=7, max_row=10, max_col=6, values_only=True):
    if txt(r[0]):
        t4s.append({"key": txt(r[0]), "label": txt(r[0]), "ff": rnd(r[2]), "df": rnd(r[3]), "nf": rnd(r[4]), "total": rnd(r[5])})
save("t5_t4000_sum", t4s)


t4 = []
for i, r in enumerate(sheet(f, "Data").iter_rows(min_row=2, max_col=16, values_only=True)):
    if not txt(r[2]):
        continue
    stock, avg, oo = num(r[11]) or 0, num(r[9]) or 0, num(r[14]) or 0
    flag = ("Stock ≤ 0 ไม่มี Order" if stock <= 0 and oo <= 0 else "Stock ≤ 0 มี Order" if stock <= 0 else
            "Stock Day < 1" if avg and stock / avg < 1 else "มีของ")
    t4.append({"key": code(r[2]), "rank": i + 1, "dept": code(r[0]), "division": div_of(r[0]), "cls": code(r[1]), "item": code(r[2]),
               "name": txt(r[3]), "status": txt(r[4]), "ordering": txt(r[5]), "reorder": txt(r[6]), "last_del": day(r[7]), "last_sold": day(r[8]),
               "avg_qty": rnd(r[9]), "avg_amt": rnd(r[10]), "stock": rnd(stock), "stock_day": round(stock / avg, 1) if avg else None,
               "on_order": rnd(oo), "promotion": txt(r[15]), "flag": flag})
save("t5_t4000", dedupe(t4))

f = find("St68 Top1000.xlsx")
ws = sheet(f, "Stock 0")
t1s = []
for r in ws.iter_rows(min_row=5, max_row=14, max_col=8, values_only=True):
    t1s.append({"key": txt(r[1]) or "Total", "section": txt(r[1]) or "Total", "sd1": rnd(r[2]), "sd1_oo": rnd(r[3]), "sd1_no": rnd(r[4]),
                "s0": rnd(r[5]), "s0_oo": rnd(r[6]), "s0_no": rnd(r[7])})
save("t5_t1000_sum", t1s)
t1 = []
for i, r in enumerate(sheet(f, "Data").iter_rows(min_row=2, max_col=25, values_only=True)):
    if not txt(r[8]):
        continue
    stock, avg, oo = num(r[20]) or 0, num(r[18]) or 0, num(r[23]) or 0
    flag = ("Stock ≤ 0 ไม่มี Order" if stock <= 0 and oo <= 0 else "Stock ≤ 0 มี Order" if stock <= 0 else
            "Stock Day < 1" if avg and stock / avg < 1 else "มีของ")
    t1.append({"key": code(r[8]), "rank": i + 1, "dept": code(r[4]), "division": div_of(r[4]), "dept_name": txt(r[5]), "cls": code(r[6]),
               "cls_name": txt(r[7]), "item": code(r[8]), "name": txt(r[9]), "supplier": txt(r[11]), "status": txt(r[13]), "reorder": txt(r[15]),
               "last_del": day(r[16]), "last_sold": day(r[17]), "avg_qty": rnd(r[18]), "avg_amt": rnd(r[19]), "stock": rnd(stock),
               "stock_day": round(stock / avg, 1) if avg else None, "on_order": rnd(oo), "promotion": txt(r[24]), "flag": flag})
save("t5_t1000", dedupe(t1))

# ---------- Top 50 FF ----------
f = find("St.68 Top50FF.xlsb")
from pyxlsb import open_workbook
with open_workbook(f) as wb:
    names = [s for s in wb.sheets if s.startswith(("Dept", "FZ"))]
t50, d50 = [], None
for s in names:
    rows = xlsb_rows(f, s)
    if len(rows) > 3 and len(rows[3]) > 2:
        d50 = d50 or day(rows[3][2])
    for r in rows[6:]:
        r = r + [None] * 24
        if num(r[1]) is None or num(r[2]) is None:
            continue
        t50.append({"key": f"{s.strip()}|{code(r[2])}", "section": s.strip(), "no": int(num(r[1])), "item": code(r[2]), "desc": txt(r[3]),
                    "qty": rnd(r[4]), "amount": rnd(r[5]), "price": rnd(r[6]), "stock": rnd(r[7]), "avg28": rnd(r[8]), "mtd": rnd(r[9]),
                    "lmtd": rnd(r[10]), "diff": rnd(r[11]), "pct": rnd(r[12]), "reorder": txt(r[19])})
save("t5_top50", dedupe(t50))
meta.append({"key": "Top 50 FF date", "value": d50})

# ---------- Sales by Category (thousand baht) ----------
f = find("ST.68 Sales By Category.xlsx")
P = ["d", "m", "y"]
F7 = ["a", "b", "var_b", "pct_b", "ly", "var_ly", "pct_ly"]


def cat_vals(vals):
    out = {}
    for p in range(3):
        for j, k in enumerate(F7):
            v = vals[p * 7 + j]
            out[f"{P[p]}_{k}"] = pct(v) if k.startswith("pct") else rnd(v)
    return out


ws = sheet(f, "Summary")
meta.append({"key": "Sales by Category date", "value": day(ws["F2"].value)})
csum = []
for r in ws.iter_rows(min_row=4, max_row=10, max_col=23, values_only=True):
    if txt(r[1]):
        k = "ALL" if txt(r[1]) == "Grand Total" else txt(r[1])
        csum.append({"key": k, "division": k, **cat_vals(r[2:23])})
save("t5_cat_div", csum)
ccat, dv = [], ""
for r in ws.iter_rows(min_row=15, max_row=60, max_col=23, values_only=True):
    if txt(r[0]) and not txt(r[0]).startswith("Total"):
        dv = txt(r[0])
    if txt(r[1]) and dv:
        ccat.append({"key": f"{dv}|{txt(r[1])}", "division": dv, "category": txt(r[1]), **cat_vals(r[2:23])})
save("t5_cat_cat", ccat)
ccls = []
for s in ["DG FOOD", "DG NONFOOD", "FRESH FOOD", "HARDLINE E", "HARDLINE H", "LIQUOR & TOBACCO"]:
    cat = ""
    for r in sheet(f, s).iter_rows(min_row=4, max_col=25, values_only=True):
        if txt(r[1]):
            cat = txt(r[1])
        if num(r[2]) is None or all(num(v) is None for v in r[4:25]):
            continue
        ccls.append({"key": f"{s}|{code(r[2])}", "division": s, "category": cat, "cls": code(r[2]), "desc": txt(r[3]), **cat_vals(r[4:25])})
save("t5_cat_cls", dedupe(ccls))

# ---------- Sale by item / Top drop customer ----------
ws = sheet(find("Sale by item.xlsx"))
save("t5_sale_item", dedupe([{"key": code(r[3]), "dept": code(r[0]), "division": div_of(r[0]), "cls": code(r[1]), "cls_name": txt(r[2]),
                               "item": code(r[3]), "desc": txt(r[4]), "cy": rnd(r[5]), "ly": rnd(r[6]), "var": rnd(r[7])}
                              for r in ws.iter_rows(min_row=4, max_col=8, values_only=True) if num(r[0]) is not None and txt(r[3])]))
ws = sheet(find("Top drop Customer by class1.xlsx"))
by_item, by_type, dv = {}, {}, ""
for r in ws.iter_rows(min_row=1, max_col=14, values_only=True):
    if txt(r[0]).upper() in ("FRESH FOOD", "DRY FOOD", "NON FOOD"):
        dv = txt(r[0]).upper(); continue
    if num(r[0]) is None or not txt(r[4]):
        continue
    cy, ly, var = num(r[9]) or 0, num(r[10]) or 0, num(r[11]) or 0
    k = code(r[4])
    o = by_item.setdefault(k, {"key": k, "division": dv, "dept": code(r[1]), "cls": code(r[2]), "cls_name": txt(r[3]), "item": k,
                                "desc": txt(r[5]), "customers": 0, "cy": 0.0, "ly": 0.0, "var": 0.0, "lost": 0})
    o["customers"] += 1; o["cy"] += cy; o["ly"] += ly; o["var"] += var; o["lost"] += cy == 0
    t = by_type.setdefault(f"{dv}|{txt(r[6])}", {"key": f"{dv}|{txt(r[6])}", "division": dv, "type": txt(r[6]) or "(ว่าง)",
                                                "customers": 0, "cy": 0.0, "ly": 0.0, "var": 0.0, "lost": 0})
    t["customers"] += 1; t["cy"] += cy; t["ly"] += ly; t["var"] += var; t["lost"] += cy == 0
for d in (by_item, by_type):
    for o in d.values():
        for f2 in ("cy", "ly", "var"):
            o[f2] = round(o[f2], 2)
save("t5_drop_item", list(by_item.values()))
save("t5_drop_type", list(by_type.values()))

# ---------- Hamper ----------
f = find("Sales_Hamper 2026.xlsm")
ws = sheet(f, "Summary")
meta.append({"key": "Hamper period", "value": txt(ws["F2"].value)})
save("t5_hamper_sup", [{"key": code(r[1]) or "Total", "supplier": txt(r[2]) or "Total", "is_total": not txt(r[2]), "stock_qty": rnd(r[3]),
                        "stock_amt": rnd(r[4]), "ly_qty": rnd(r[5]), "ly_amt": rnd(r[6]), "cy_qty": rnd(r[7]), "cy_amt": rnd(r[8])}
                       for r in ws.iter_rows(min_row=6, max_row=20, max_col=9, values_only=True) if num(r[3]) is not None])
hi = []
for r in sheet(f, "SalesLYCY").iter_rows(min_row=5, max_col=22, values_only=True):
    if not txt(r[5]):
        continue
    hi.append({"key": code(r[5]), "dept": code(r[1]), "cls": code(r[3]), "item": code(r[5]), "desc": txt(r[6]), "ly_qty": rnd(r[7]),
               "ly_amt": rnd(r[8]), "ly_profit": rnd(r[9]), "cy_qty": rnd(r[11]), "cy_amt": rnd(r[12]), "cy_profit": rnd(r[13]),
               "var": rnd(r[15]), "soh": rnd(r[19]), "soh_cost": rnd(r[20]), "on_order": rnd(r[21])})
save("t5_hamper_items", dedupe(hi))
ws = sheet(f, "Daily_Sales")
rows = list(ws.iter_rows(min_row=4, max_col=40, values_only=True))
dates = rows[0]
hd = []
for c in range(4, 40, 2):
    dd = day(dates[c])
    if not dd:
        continue
    q = sum(num(r[c]) or 0 for r in rows[2:] if txt(r[2]))
    a = sum(num(r[c + 1]) or 0 for r in rows[2:] if txt(r[2]))
    hd.append({"key": dd, "date": dd, "qty": round(q, 2), "amount": round(a, 2)})
save("t5_hamper_day", hd)

# ---------- Shrinkage by code MTD ----------
ws = sheet(find("Shrinkage_by_code_by_Div_MTD.xlsx"))
CODES = ["c10", "c11", "c20", "c21", "c52", "c53", "c60", "c61", "tc41"]
sc, dv = [], ""
for r in ws.iter_rows(min_row=1, max_col=17, values_only=True):
    if txt(r[0]).upper() in ("FRESH FOOD", "DRY FOOD", "NON FOOD"):
        dv = txt(r[0]).upper(); continue
    if num(r[0]) is None or not txt(r[2]):
        continue
    rec = {"key": code(r[2]), "division": dv, "dept": code(r[0]), "cls": code(r[1]), "item": code(r[2]), "desc": txt(r[3])}
    rec.update({c: rnd(v) for c, v in zip(CODES, r[4:13])})
    rec.update({"qty": rnd(r[13]), "shrink": rnd(r[14]), "sale": rnd(r[15]), "pct": pct(r[16], 3)})
    sc.append(rec)
save("t5_shk_code", dedupe(sc))

# ---------- Last Sales 3 Day ----------
f = find("ST.68 Last Sales (3 Day)*.xlsx")
save("t5_last3_sum", [{"key": DIVN.get(txt(r[1]), "ALL"), "division": DIVN.get(txt(r[1]), "ALL"), "items": rnd(r[2]), "s0": rnd(r[3]),
                       "s5": rnd(r[4]), "s10": rnd(r[5]), "on_order": rnd(r[6])}
                      for r in sheet(f, "Summary").iter_rows(min_row=7, max_row=10, max_col=7, values_only=True) if txt(r[1])])
ws = openpyxl.load_workbook(f, read_only=True, data_only=True).worksheets[1]
save("t5_last3", dedupe([{"key": code(r[2]), "dept": code(r[0]), "division": div_of(r[0]), "cls": code(r[1]), "item": code(r[2]),
                          "desc": txt(r[3]), "last_sold": day(r[4]), "soh": rnd(r[5]), "on_order": rnd(r[6]),
                          "band": "Stock = 0" if (num(r[5]) or 0) <= 0 else "Stock < 5" if num(r[5]) < 5 else "Stock < 10" if num(r[5]) < 10 else "Stock ≥ 10"}
                         for r in ws.iter_rows(min_row=2, max_col=9, values_only=True) if num(r[0]) is not None and txt(r[2])]))

# ---------- Negative stock ----------
f = find("St68 Negative stock.xlsx")
ws = sheet(f, "สรุป")
meta.append({"key": "Negative stock date", "value": day(ws["D3"].value)})
save("t5_neg_sum", [{"key": txt(r[1]), "section": txt(r[1]), "count": rnd(r[2]), "adjusted": rnd(r[3]), "open": rnd(r[4])}
                    for r in ws.iter_rows(min_row=6, max_row=15, max_col=5, values_only=True) if txt(r[1])])
neg = []
for r in sheet(f, "Nagative").iter_rows(min_row=2, max_col=27, values_only=True):
    a = num(r[9])
    if a is None or a >= 0 or not txt(r[4]):
        continue
    neg.append({"key": code(r[4]), "buyer": txt(r[1]), "dept": code(r[2]), "division": div_of(r[2]), "cls": code(r[3]), "item": code(r[4]),
                "desc": txt(r[5]), "last_sale": day(r[6]), "last_del": day(r[7]), "price": rnd(r[8]), "asoh": rnd(a), "value": rnd(a * (num(r[8]) or 0)),
                "on_order": rnd(r[11]), "avg28": rnd(r[13]), "source": txt(r[18]), "reorder": txt(r[19]), "status": txt(r[22])})
save("t5_neg_items", dedupe(neg))
adj = {}
for r in sheet(f, "Adj").iter_rows(min_row=2, max_col=15, values_only=True):
    if not txt(r[7]):
        continue
    k = f"{txt(r[2])}|{txt(r[14])}"
    o = adj.setdefault(k, {"key": k, "reason": txt(r[2]) or "(ว่าง)", "status": txt(r[14]) or "(ว่าง)", "lines": 0, "qty": 0.0, "value": 0.0})
    o["lines"] += 1; o["qty"] = round(o["qty"] + (num(r[12]) or 0), 2); o["value"] = round(o["value"] + (num(r[13]) or 0), 2)
# the page does not show adjustments yet; they are summed only so the sheet is checked to read cleanly
print(f"t5_neg_adj: {len(adj)} groups (not saved)")

# ---------- Gone Forever / Discontinue ----------


def gone(f, prefix, sheets):
    ws = sheet(f, "Summary")
    sm = []
    for r in ws.iter_rows(min_row=7, max_row=12, max_col=11, values_only=True):
        if txt(r[0]):
            u = txt(r[0]).upper()
            k = ("ALL" if u == "TOTAL" else "NON FOOD" if u.startswith("NON") else "DRY FOOD" if u == "DRY FOOD" else
                 "FRESH FOOD" if u.startswith("FRESH") else txt(r[0]))
            sm.append({"key": k, "division": k, "start_qty": rnd(r[1]), "start_amt": rnd(r[2]), "now_qty": rnd(r[3]), "now_amt": rnd(r[4]),
                       "var_qty": rnd(r[5]), "var_amt": rnd(r[6]), "pct_qty": pct(r[7]), "pct_amt": pct(r[8]), "sales_qty": rnd(r[9]), "sales_amt": rnd(r[10])})
    save(prefix + "_sum", sm)
    items, daily = {}, defaultdict(lambda: [0.0, 0.0])
    for s in sheets:
        rows = list(sheet(f, s).iter_rows(min_row=4, max_col=71, values_only=True))
        dates = rows[0]
        for r in rows[1:]:
            if not txt(r[2]) or num(r[2]) is None:
                continue
            k = code(r[2])
            if k in items:
                continue
            q = a = 0.0; last = None
            for c in range(7, 69, 2):
                dq, da = num(r[c]) or 0, num(r[c + 1]) or 0
                dd = day(dates[c])
                if dd and (dq or da):
                    q += dq; a += da; last = dd
                    daily[dd][0] += dq; daily[dd][1] += da
            items[k] = {"key": k, "div_code": txt(r[0]), "division": div_of(r[1]), "dept": code(r[1]), "item": k, "desc": txt(r[3]),
                        "on_order": rnd(r[4]), "stock_qty": rnd(r[5]), "stock_amt": rnd(r[6]), "sales_qty": round(q, 2),
                        "sales_amt": round(a, 2), "last_sale": last}
    save(prefix + "_items", list(items.values()))
    save(prefix + "_day", [{"key": d, "date": d, "qty": round(v[0], 2), "amount": round(v[1], 2)} for d, v in sorted(daily.items())])


gone(find("ST68 Gone Forever Sales*.xlsx"), "t5_gone", ["Gone Forever FF", "Gone Forever DF", "Gone Forever NF"])
gone(find("ST68 Sales Discontinue.xlsx"), "t5_disc", ["Discontinue"])

# ---------- Soft-serve & Rotisserie ----------
ws = sheet(find("St.68 Soft-servers Rotisserie*.xlsx"), "Sum")
save("t5_soft", [{"key": day(r[0]), "date": day(r[0]), "label": txt(r[1]), "soft_qty": rnd(r[2]), "soft_amt": rnd(r[3]), "chk_qty": rnd(r[4]),
                  "chk_amt": rnd(r[5]), "durian_qty": rnd(r[6]), "durian_amt": rnd(r[7]), "qty": rnd(r[8]), "amount": rnd(r[9])}
                 for r in ws.iter_rows(min_row=6, max_col=10, values_only=True) if day(r[0]) and num(r[8]) is not None])
save("t5_meta", meta)

"""Classes that have sales in ST.68 Sales By Category but sit in none of its category sheets (the file's own
"Check Diff"). Both extractors add them as a group "ไม่มี Category" so category totals match the store total.

unmapped_classes(path) -> list of dicts, money in thousand baht like the file:
  dept, dept_name, cls, desc, store_div, and d_/m_/y_ + a (actual), b (budget), ly (last year).
Today comes from the Daily Sales / Daily Sales LY sheets, MTD and YTD from the Data sheet.
"""
import warnings
import openpyxl

warnings.filterwarnings("ignore")
CATEGORY_SHEETS = ["DG FOOD", "DG NONFOOD", "FRESH FOOD", "HARDLINE E", "HARDLINE H", "LIQUOR & TOBACCO"]
GROUP = "ไม่มี Category"
LABEL = "Class ที่ยังไม่มี Category"


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _code(v):
    n = _num(v)
    return str(int(n)) if n is not None and n == int(n) else ("" if v is None else str(v).strip())


def _rows(ws, first, cols):
    """Class rows with division / department carried down from the row that names them."""
    div = dept = dname = None
    for r in ws.iter_rows(min_row=first, max_col=cols, values_only=True):
        if r[0]:
            div = str(r[0]).strip().upper()
        if r[1] is not None:
            dept, dname = _code(r[1]), str(r[2] or "").strip()
        if _num(r[3]) is None:
            continue
        yield div, dept, dname, _code(r[3]), str(r[4] or "").strip(), r


def unmapped_classes(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    mapped = set()
    for s in CATEGORY_SHEETS:
        if s in wb.sheetnames:
            for r in wb[s].iter_rows(min_row=4, max_col=3, values_only=True):
                if _num(r[2]) is not None:
                    mapped.add(_code(r[2]))
    day, day_ly = {}, {}
    if "Daily Sales" in wb.sheetnames:
        for *_, cls, _d, r in _rows(wb["Daily Sales"], 6, 7):
            day[cls] = (_num(r[5]) or 0, _num(r[6]) or 0)
    if "Daily Sales LY" in wb.sheetnames:
        for *_, cls, _d, r in _rows(wb["Daily Sales LY"], 7, 6):
            day_ly[cls] = _num(r[5]) or 0
    out = []
    for div, dept, dname, cls, desc, r in _rows(wb["Data"], 5, 11):
        if cls in mapped:
            continue
        vals = [_num(x) or 0 for x in r[5:11]]
        da, db = day.get(cls, (0, 0))
        if not any(vals) and not da:
            continue
        out.append({"dept": dept, "dept_name": dname, "cls": cls, "desc": desc, "store_div": div,
                    "d_a": da, "d_b": db, "d_ly": day_ly.get(cls, 0),
                    "m_a": vals[0], "m_b": vals[1], "m_ly": vals[2], "y_a": vals[3], "y_b": vals[4], "y_ly": vals[5]})
    return out


if __name__ == "__main__":
    import glob, sys
    sys.stdout.reconfigure(encoding="utf-8")
    f = sorted(glob.glob("**/ST.68 Sales By Category.xlsx", recursive=True))[-1]
    rows = unmapped_classes(f)
    for x in rows:
        print(x["store_div"], x["dept"], x["cls"], x["desc"], round(x["d_a"], 2), round(x["m_a"], 2), round(x["y_a"], 2))
    print("MTD total:", round(sum(x["m_a"] for x in rows), 2), " today:", round(sum(x["d_a"] for x in rows), 2))

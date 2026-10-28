"""Topic 3 of the Store 68 Daily Board: the four Focus files (Top 1000 O2O, All FF sale & profit, Code 20 shrinkage,
Shrinkage by department / class / item). Reads them from anywhere under this folder and writes JSON datasets.

E-mail addresses (Summary Daily FF column B, the Set Email sheet) are never read out. Percentages are written
as percent units (2.5 = 2.5%). The 168k-row Item List is summed up per department and status, not copied.
Usage: python build_topic3.py [out_folder]
"""
import glob, json, os, re, sys, warnings
from collections import defaultdict
from datetime import datetime, date
import openpyxl

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "data_t3")
os.makedirs(OUT, exist_ok=True)
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def find(pattern):
    hits = sorted(glob.glob(os.path.join(HERE, "**", pattern), recursive=True), key=os.path.getmtime)
    if not hits:
        sys.exit(f"{pattern} not found under {HERE}")
    print("reading", hits[-1])
    return hits[-1]


def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("%", "").replace("~", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def rnd(v, d=2):
    v = num(v)
    return None if v is None else round(v, d)


def pct(v, scale=100, d=2):
    """Percent units; fractions in the sheet are multiplied by 100."""
    v = num(v)
    return None if v is None else round(v * scale, d)


def txt(v):
    return "" if v is None else str(v).strip()


def code(v):
    n = num(v)
    return str(int(n)) if n is not None and n == int(n) else txt(v)


def day(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    m = re.match(r"(\d{2})-(\d{2})-(\d{4})", txt(v))
    if m and m.group(3) > "1999":
        return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", txt(v))
    return f"{m.group(3)}-{int(m.group(1)):02d}-{int(m.group(2)):02d}" if m else None


def div_of(dept):
    n = num(dept) or 0
    return "FRESH FOOD" if n <= 5 or n in (19, 21) else "DRY FOOD" if n <= 17 or n == 22 else "NON FOOD"


def save(name, rows):
    with open(os.path.join(OUT, name + ".json"), "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, separators=(",", ":"))
    print(f"{name}: {len(rows)} rows")


meta = []

# ---------- Focus Top 1000 O2O ----------
wb = openpyxl.load_workbook(find("Focus Top 1000*O2O*.xlsx"), read_only=True, data_only=True)
DIVS = {"FF": "FRESH FOOD", "DF": "DRY FOOD", "NF": "NON FOOD"}
summ = []
for r in wb["Summary"].iter_rows(min_row=5, max_row=12, max_col=8, values_only=True):
    d = txt(r[1])
    if not d:
        continue
    summ.append({"division": DIVS.get(d, "ALL"), "code": d, "on_order": rnd(r[2]), "in_transit": rnd(r[3]),
                 "zero_stock": rnd(r[4]), "zero_on_order": rnd(r[5]), "no_order": rnd(r[6]), "negative": rnd(r[7])})
save("t3_top_sum", summ)

ra = {}
ws = wb["RA"]
for r in ws.iter_rows(min_row=3, max_col=29, values_only=True):
    for i, dv in ((0, "FRESH FOOD"), (10, "DRY FOOD"), (20, "NON FOOD")):
        b = r[i:i + 9]
        if num(b[0]) is None or not txt(b[3]):
            continue
        ra[(dv, code(b[3]))] = {"sales_rank": int(num(b[0])), "sales_qty": rnd(b[5]), "sales": rnd(b[6]),
                                "profit": rnd(b[7]), "profit_pct": rnd(b[8])}

top, data_date = [], None
for sh, dv in DIVS.items():
    for r in wb[sh].iter_rows(min_row=4, max_col=13, values_only=True):
        if num(r[0]) is None or not txt(r[3]):
            continue
        item = code(r[3])
        stock_txt = txt(r[5])
        stock, oo, it = num(r[5]), num(r[6]) or 0, num(r[7]) or 0
        if stock is None:
            stock = 0.0
        flag = ("Stock ติดลบ" if stock < 0 else
                "หมด มี Order" if stock <= 0 and (oo > 0 or it > 0) else
                "หมด ไม่มี Order" if stock <= 0 else
                "มีของ มี Order" if oo > 0 else "มีของ In Transit" if it > 0 else "มีของ ไม่มี Order")
        last_sale = day(r[11])
        if last_sale and (data_date is None or last_sale > data_date):
            data_date = last_sale
        rec = {"key": f"{dv}|{item}", "division": dv, "rank": int(num(r[0])), "dept": code(r[1]), "cls": code(r[2]),
               "item": item, "name": txt(r[4]), "stock": round(stock, 2), "approx": stock_txt.startswith("~"),
               "on_order": rnd(oo), "in_transit": rnd(it), "stock_day": rnd(r[8], 0), "avg28": rnd(r[9]),
               "status": txt(r[10]), "last_sale": last_sale, "last_del": day(r[12]), "flag": flag}
        rec.update(ra.get((dv, item), {}))
        top.append(rec)
save("t3_top_items", top)
meta.append({"metric": "Top 1000 date", "value": data_date})

# Item List: every article of the store, summed per department and by status
ws = wb["Item List"]
agg = defaultdict(lambda: defaultdict(float))
stat = defaultdict(lambda: defaultdict(float))
hdr_seen = False
for r in ws.iter_rows(min_row=1, max_col=28, values_only=True):
    if not hdr_seen:
        hdr_seen = txt(r[1]) == "Supplier" and txt(r[3]).startswith("Dept")
        continue
    if num(r[3]) is None or not txt(r[5]):
        continue
    dept = code(r[3])
    price, usoh = num(r[9]) or 0, num(r[11]) or 0
    status = txt(r[23]) or "(ว่าง)"
    aging = txt(r[24]) or "(ว่าง)"
    a = agg[dept]
    a["items"] += 1
    a["active"] += status == "Active"
    a["with_stock"] += usoh > 0
    a["active_zero"] += status == "Active" and usoh <= 0
    a["negative"] += usoh < 0
    a["stock_value"] += usoh * price if usoh > 0 else 0
    a["on_order"] += (num(r[12]) or 0) > 0
    a["mmail"] += bool(txt(r[27]))
    a["sales_cm"] += num(r[16]) or 0
    a["sales_lm"] += num(r[15]) or 0
    a["sales_ly"] += num(r[18]) or 0
    for dim, val in (("Item Status", status), ("Aging Status", aging), ("Source", txt(r[19]) or "(ว่าง)"),
                     ("Reorder Way", txt(r[20]) or "(ว่าง)")):
        s = stat[(div_of(dept), dim, val)]
        s["items"] += 1
        s["with_stock"] += usoh > 0
        s["stock_value"] += usoh * price if usoh > 0 else 0
il_dept = []
for d, a in sorted(agg.items(), key=lambda x: num(x[0]) or 0):
    il_dept.append({"key": d, "dept": d, "division": div_of(d), **{k: round(v, 2) for k, v in a.items()}})
save("t3_il_dept", il_dept)
save("t3_il_status", [{"key": f"{k[0]}|{k[1]}|{k[2]}", "division": k[0], "dim": k[1], "value": k[2],
                       **{f: round(v, 2) for f, v in s.items()}} for k, s in stat.items()])

# ---------- Focus_All_FF.XLSM ----------
wb = openpyxl.load_workbook(find("Focus_All_FF.XLSM"), read_only=True, data_only=True, keep_vba=False)
SUMF = ["sale", "profit", "profit_pct", "em_sale", "em_profit", "em_pct", "c52", "c52_pct", "tt_sale",
        "tt_profit", "tt_pct", "loss_margin"]
ff_sum = []
for sh, per in (("Summary Daily FF", "Daily"), ("Summary MTD FF", "MTD")):
    for r in wb[sh].iter_rows(min_row=5, max_col=17, values_only=True):
        dep = num(r[3])
        if dep is None:
            if ff_sum and ff_sum[-1]["period"] == per:
                break
            continue
        if dep >= 100:
            break
        rec = {"key": f"{per}|{code(r[3])}", "period": per, "dept": code(r[3]), "desc": txt(r[4])}
        for f, v in zip(SUMF, r[5:17]):
            rec[f] = pct(v) if f.endswith("pct") else rnd(v)
        ff_sum.append(rec)
save("t3_ff_sum", ff_sum)
ws = wb["Summary Daily FF"]
meta.append({"metric": "All FF date", "value": day(ws["O3"].value)})

ff_items = []
for sh, per in (("Sum_D", "Daily"), ("Sum_M", "MTD")):
    for r in wb[sh].iter_rows(min_row=5, max_col=25, values_only=True):
        if num(r[3]) is None:
            continue
        rec = {"key": f"{per}|{code(r[3])}", "period": per, "dept": code(r[1]), "cls": code(r[2]), "item": code(r[3]),
               "desc": txt(r[4]), "dms": rnd(r[5]), "stock": rnd(r[6]), "aging": rnd(r[7], 0),
               "sale": rnd(r[8]), "profit": rnd(r[9]), "profit_pct": rnd(r[10]), "em_sale": rnd(r[11]),
               "em_profit": rnd(r[12]), "em_pct": rnd(r[13]), "c52": rnd(r[14]), "tt_sale": rnd(r[16]),
               "tt_profit": rnd(r[17]), "tt_pct": pct(r[18]), "em_item": "" if txt(r[19]) == "NO" else code(r[19])}
        if any(rec[k] for k in ("sale", "em_sale", "c52", "stock")):
            ff_items.append(rec)
save("t3_ff_items", ff_items)

c52 = []
for sh, per, start in (("Code_52_Daily", "Daily", 4), ("Code_52_MTD", "MTD", 5)):
    for r in wb[sh].iter_rows(min_row=start, max_col=8, values_only=True):
        if num(r[1]) is None:
            continue
        c52.append({"key": f"{per}|{code(r[1])}", "period": per, "item": code(r[1]), "dept": code(r[2]),
                    "cls": code(r[3]), "desc": txt(r[4]), "qty": rnd(r[5]), "cost": rnd(r[6])})
merged = {}
for c in c52:
    m = merged.setdefault(c["key"], dict(c, qty=0, cost=0))
    m["qty"] = round(m["qty"] + (c["qty"] or 0), 3)
    m["cost"] = round(m["cost"] + (c["cost"] or 0), 2)
save("t3_ff_c52", list(merged.values()))

# ---------- Focus Code 20 By Item ----------
wb = openpyxl.load_workbook(find("Focus Code 20 By Item.xlsx"), read_only=True, data_only=True)
c20m = []
for r in wb["Summary"].iter_rows(min_row=4, max_row=15, max_col=8, values_only=True):
    if txt(r[1]) in MONTHS:
        c20m.append({"key": txt(r[1]), "month": txt(r[1]), "m": MONTHS.index(txt(r[1])) + 1,
                     "shrink_ly": rnd(r[2]), "sale_ly": rnd(r[3]), "pct_ly": pct(r[4], d=3),
                     "shrink": rnd(r[5]), "sale": rnd(r[6]), "pct": pct(r[7], d=3)})
save("t3_c20_month", c20m)
c20i = []
for mo in MONTHS:
    if mo not in wb.sheetnames:
        continue
    ws = wb[mo]
    for r in ws.iter_rows(min_row=9, max_col=17, values_only=True):
        for off, yr in ((1, 2025), (9, 2026)):
            b = r[off:off + 8]
            if num(b[2]) is None:
                continue
            c20i.append({"key": f"{yr}|{MONTHS.index(mo) + 1}|{code(b[2])}", "year": yr, "m": MONTHS.index(mo) + 1,
                         "dept": code(b[0]), "cls": code(b[1]), "item": code(b[2]), "desc": txt(b[3]),
                         "qty": rnd(b[4]), "shrink": rnd(b[5]), "sale": rnd(b[6])})
merged = {}
for c in c20i:
    m = merged.get(c["key"])
    if m:
        for f in ("qty", "shrink", "sale"):
            m[f] = round((m[f] or 0) + (c[f] or 0), 2)
    else:
        merged[c["key"]] = c
# item names go to their own small dataset so the 30k monthly rows stay light
names = {}
for c in merged.values():
    names.setdefault(c.pop("item"), c.pop("desc"))
    c["item"] = c["key"].rsplit("|", 1)[1]
save("t3_c20_items", list(merged.values()))
save("t3_c20_names", [{"key": k, "item": k, "desc": v} for k, v in names.items()])

# ---------- Focus Shrinkage ----------
wb = openpyxl.load_workbook(find("Focus Shrinkage.xlsx"), read_only=True, data_only=True)
shk, div = [], ""
last_dept = {"dept": "", "name": "", "div": ""}
for r in wb["Sum"].iter_rows(min_row=4, max_col=14, values_only=True):
    if txt(r[0]):
        div = txt(r[0]).upper()
    if num(r[1]) is None:
        if txt(r[0]) or txt(r[2]):
            label = txt(r[0]) or txt(r[2])
            shk.append({"key": "T|" + label, "dept": "", "desc": label, "division": "ALL" if "TOTAL" in label.upper() and "FOOD" not in label.upper() else div, "is_total": True,
                        **dict(zip(["d_sales", "d_shrink", "d_pct", "m_sales", "m_shrink", "m_pct", "m_budget",
                                    "y_sales", "y_shrink", "y_pct", "y_budget"],
                                   [rnd(v) for v in r[3:14]]))})
        continue
    shk.append({"key": code(r[1]), "dept": code(r[1]), "desc": txt(r[2]), "division": div, "is_total": False,
                **dict(zip(["d_sales", "d_shrink", "d_pct", "m_sales", "m_shrink", "m_pct", "m_budget",
                            "y_sales", "y_shrink", "y_pct", "y_budget"], [rnd(v) for v in r[3:14]]))})
save("t3_shk_dept", shk)
meta.append({"metric": "Shrinkage date", "value": day(wb["Sum"]["A2"].value)})

COMP = ["sales", "md_sales", "adj", "count", "damage", "stolen", "markdown", "donation", "shrink", "budget"]
cls = []
for r in wb["data"].iter_rows(min_row=6, max_col=45, values_only=True):
    if num(r[3]) is None:
        continue
    if num(r[1]) is not None:
        last_dept = {"dept": code(r[1]), "name": txt(r[2]), "div": txt(r[0]).upper() or div_of(r[1])}
    rec = {"key": f"{last_dept['dept']}|{code(r[3])}", "division": last_dept["div"], "dept": last_dept["dept"],
           "dept_name": txt(r[2]) or last_dept["name"], "cls": code(r[3]), "cls_name": txt(r[4])}
    for f, v in zip(COMP, r[5:15]):
        rec["d_" + f] = rnd(v)
    rec["d_pct"], rec["d_budget_pct"] = pct(r[15], d=3), pct(r[16], d=3)
    for f, v in zip(COMP + ["ly"], r[17:28]):
        rec["m_" + f] = rnd(v)
    rec["m_pct"], rec["m_budget_pct"], rec["m_var_ly"] = pct(r[28], d=3), pct(r[29], d=3), rnd(r[30])
    for f, v in zip(COMP + ["ly"], r[31:42]):
        rec["y_" + f] = rnd(v)
    rec["y_pct"], rec["y_var_budget"], rec["y_var_ly"] = pct(r[42], d=3), rnd(r[43]), rnd(r[44])
    cls.append(rec)
save("t3_shk_class", cls)

items = []
for r in wb["shk by item"].iter_rows(min_row=6, max_col=8, values_only=True):
    if num(r[0]) is None or num(r[2]) is None:
        continue
    items.append({"key": f"{code(r[2])}|{txt(r[7])}", "dept": code(r[0]), "division": div_of(r[0]), "cls": code(r[1]),
                  "item": code(r[2]), "desc": txt(r[3]), "qty": rnd(r[4]), "cost": rnd(r[5]), "cost_mtd": rnd(r[6]),
                  "reason": code(r[7])})
merged = {}
for c in items:
    m = merged.get(c["key"])
    if m:
        for f in ("qty", "cost", "cost_mtd"):
            m[f] = round((m[f] or 0) + (c[f] or 0), 2)
    else:
        merged[c["key"]] = c
save("t3_shk_items", list(merged.values()))
save("t3_meta", [dict(m, key=m["metric"]) for m in meta])
